"""Lectura del Excel oficial, selección por periodo, detección de cambios e instalación."""

import shutil

import pytest
from openpyxl import load_workbook

from fiscal_es.core import diseno, disenos_aeat
from fiscal_es.core.periodo import Trimestre

XLSX = diseno.DIR_EMPAQUETADO / "DR303e26v101.xlsx"


def test_lee_excel_oficial():
    d = diseno.leer_xlsx(XLSX)
    assert {"00", "01", "02", "03", "04", "05", "DID"} <= set(d.paginas)
    assert (d.paginas["01"].total, d.paginas["03"].total, d.paginas["DID"].total) == (1581, 1017, 823)
    c = {x.casilla: x for x in d.paginas["01"].campos if x.casilla}
    assert (c["07"].pos, c["07"].lon, c["07"].tipo) == (326, 17, "Num")
    assert c["08"].constante == "02100"
    assert c["46"].tipo == "N"


def test_seleccion_por_periodo(tmp_path):
    assert diseno.seleccionar(Trimestre(2026, 1), tmp_path).entrada.archivo == "DR303e26v101.xlsx"
    assert diseno.seleccionar(Trimestre(2028, 2), tmp_path).entrada.archivo == "DR303e26v101.xlsx"
    with pytest.raises(LookupError):
        diseno.seleccionar(Trimestre(2025, 4), tmp_path)


def test_compatible_con_el_vigente():
    assert disenos_aeat.compatibilidad(diseno.leer_xlsx(XLSX)) == {"compatible": True, "problemas": []}


def _modificar(tmp_path, nombre, cambio):
    destino = tmp_path / nombre
    shutil.copyfile(XLSX, destino)
    # data_only: el Excel de la AEAT usa fórmulas para posiciones; se guardan los valores.
    wb = load_workbook(destino, data_only=True)
    cambio(wb)
    wb.save(destino)
    return destino


def test_detecta_campo_nuevo_desconocido(tmp_path):
    """Simula un diseño nuevo de la AEAT que añade un campo obligatorio que no sabemos rellenar."""

    def cambio(wb):
        ws = wb["DP30301"]
        for row in ws.iter_rows():
            if row[4].value and "Reservado para la AEAT" == str(row[4].value).strip():
                row[4].value = "Identificación (1) - Nueva marca inventada"
                break

    nuevo = diseno.leer_xlsx(_modificar(tmp_path, "DR303e27v100.xlsx", cambio))
    comp = disenos_aeat.compatibilidad(nuevo)
    assert not comp["compatible"] and "Nueva marca inventada" in comp["problemas"][0]
    dif = diseno.diferencias(diseno.leer_xlsx(XLSX), nuevo)
    assert "Identificación (1) - Nueva marca inventada" in dif["paginas"]["01"]["anadidos"]


def test_detecta_cambio_de_tipo(tmp_path):
    def cambio(wb):
        ws = wb["DP30303"]
        for row in ws.iter_rows():
            if row[4].value and "[59]" in str(row[4].value):
                row[3].value = "Num"  # antes 'N' (con signo)
                break

    nuevo = diseno.leer_xlsx(_modificar(tmp_path, "DR303e27v101.xlsx", cambio))
    dif = diseno.diferencias(diseno.leer_xlsx(XLSX), nuevo)
    assert not dif["iguales"]
    assert dif["paginas"]["03"]["modificados"][0]["cambios"]["tipo"] == ("N", "Num")


def test_posiciones_incoherentes(tmp_path):
    def cambio(wb):
        ws = wb["DP30303"]
        for row in ws.iter_rows():
            if row[4].value and "[59]" in str(row[4].value):
                row[2].value = 18  # longitud distinta sin mover el resto
                break

    with pytest.raises(ValueError, match="se esperaba"):
        diseno.leer_xlsx(_modificar(tmp_path, "DR303e27v102.xlsx", cambio))


def test_instalar_local_y_seleccion(tmp_path):
    copia = tmp_path / "DR303e27v100.xlsx"
    shutil.copyfile(XLSX, copia)
    dir_usuario = tmp_path / "disenos"
    res = disenos_aeat.instalar(dir_usuario, ruta_local=copia)
    assert res["desde"] == "2027-1T" and res["compatibilidad"]["compatible"]
    assert diseno.seleccionar(Trimestre(2027, 1), dir_usuario).origen == "usuario"
    assert diseno.seleccionar(Trimestre(2026, 4), dir_usuario).origen == "paquete"


def test_hash_alterado(tmp_path):
    dir_usuario = tmp_path / "disenos"
    copia = tmp_path / "DR303e27v100.xlsx"
    shutil.copyfile(XLSX, copia)
    res = disenos_aeat.instalar(dir_usuario, ruta_local=copia)
    (dir_usuario / res["instalado"]).write_bytes(b"alterado")
    with pytest.raises(ValueError, match="hash"):
        diseno.cargar(diseno.seleccionar(Trimestre(2027, 1), dir_usuario))


def test_parsear_pagina_aeat():
    html = """<li><a href="/static_files/Sede/Disenyo_registro/DR_300_399/archivos_26/DR303e26v101.xlsx"
      title="x">303 - Ejercicio 2026 y siguientes (actualizado 28/01/26)<span>(378 KB - xlsx)</span></a></li>
      <li><a href="/otro/DR390e26v100.xlsx">390</a></li>"""
    pubs = disenos_aeat.parsear_pagina(html)
    assert len(pubs) == 1
    p = pubs[0]
    assert p.url.startswith("https://sede.agenciatributaria.gob.es/static_files/")
    assert (p.ejercicio_desde, p.actualizado.isoformat(), p.desde) == (2026, "2026-01-28", "2026-1T")


def test_excel_sin_valores_calculados(tmp_path):
    """Si el Excel llega con fórmulas sin valor calculado, las posiciones se deducen igual."""
    destino = tmp_path / "DR303e26v101-sin-cache.xlsx"
    shutil.copyfile(XLSX, destino)
    wb = load_workbook(destino)  # conserva fórmulas; al guardar se pierden los valores calculados
    wb.save(destino)
    original, sin_cache = diseno.leer_xlsx(XLSX), diseno.leer_xlsx(destino)
    for cod in ("00", "01", "03", "DID"):
        a = [(c.num, c.pos, c.lon, c.casilla) for c in original.paginas[cod].campos]
        b = [(c.num, c.pos, c.lon, c.casilla) for c in sin_cache.paginas[cod].campos]
        assert a == b, cod
        assert original.paginas[cod].total == sin_cache.paginas[cod].total
