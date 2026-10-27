"""Diseño de registro del 303 leído directamente del Excel oficial de la AEAT.

No hay posiciones escritas a mano: todo sale del .xlsx (`DR303eAAvNNN.xlsx`).
Los diseños se versionan con un manifiesto que indica desde qué periodo aplica
cada uno. Hay dos manifiestos: el que viene con el paquete (`fiscal_es/disenos`)
y el del usuario (`FISCAL_DATA_DIR/disenos`), donde se instalan los nuevos.
"""

from __future__ import annotations

import hashlib
import json
import re
from dataclasses import asdict, dataclass, field
from functools import lru_cache
from importlib import resources
from pathlib import Path
from typing import Any

from .periodo import Trimestre

DIR_EMPAQUETADO = Path(str(resources.files("fiscal_es") / "disenos"))

_RE_CONSTANTE = re.compile(r'^(?:Obligatorio\s*)?(?:Constante\s*)?"([^"]*)"\s*(?:\.?\s*Nota\s*\d+\s*)*\.?\s*$', re.I)
_RE_CASILLA = re.compile(r"\[(\d+)\]\s*\.?\s*$")


@dataclass
class Campo:
    num: int
    pos: int | None
    lon: int | None  # None = longitud variable (contenido del fichero en la cabecera)
    tipo: str  # A | An | Num | N
    descripcion: str
    validacion: str
    contenido: str
    casilla: str | None
    constante: str | None

    @property
    def decimales(self) -> int:
        texto = f"{self.validacion} {self.contenido}".lower()
        return 2 if "decimales" in texto else 0

    @property
    def reservado(self) -> bool:
        return self.descripcion.lower().startswith("reservado")


@dataclass
class Pagina:
    hoja: str
    codigo: str  # '00' cabecera, '01'..'05', 'DID'
    campos: list[Campo]
    total: int | None
    notas: list[str] = field(default_factory=list)


@dataclass
class Diseno:
    archivo: str
    sha256: str
    paginas: dict[str, Pagina]
    modelo: str = "303"

    @property
    def nombre(self) -> str:
        return Path(self.archivo).stem


def _txt(v: Any) -> str:
    return "" if v is None else str(v).replace("\n", " ").strip()


def _int(v: Any) -> int | None:
    try:
        return int(str(v).strip())
    except (TypeError, ValueError):
        return None


# Páginas que tiene que traer el Excel de cada modelo para considerarlo válido.
PAGINAS_REQUERIDAS = {"303": {"00", "01", "03"}, "390": {"00", "01", "02", "06"}}


def _codigo_pagina(campos: list[Campo]) -> str | None:
    """Código de página a partir del propio diseño, no del nombre de la hoja.

    El campo "Página" lleva una constante como "01000", "02B00" o "DID00" -> "01", "02B", "DID".
    La cabecera no tiene ese campo, pero sí el "Contenido del fichero" de longitud variable -> "00".
    """
    for c in campos:
        if re.fullmatch(r"P[aá]gina\.?", c.descripcion.strip(), re.I) and c.constante:
            k = c.constante.strip()
            return k[:-3] if k.endswith("000") else k[:-2] if k.endswith("00") else k
    if any(c.lon is None and re.match(r"Contenido del fichero", c.descripcion, re.I) for c in campos):
        return "00"
    return None


def _modelo(campos: list[Campo]) -> str | None:
    for c in campos:
        if re.fullmatch(r"Modelo\.?", c.descripcion.strip(), re.I) and c.constante:
            return c.constante.strip()
    return None


def sha256_de(path: Path) -> str:
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def leer_xlsx(path: Path, modelo: str | None = None) -> Diseno:
    """Lee el Excel oficial (303, 390…). Lanza ValueError si la estructura no es la esperada.

    Si se indica `modelo`, se comprueba que el Excel sea de ese modelo.
    """
    from openpyxl import load_workbook

    path = Path(path)
    wb = load_workbook(path, read_only=True, data_only=True)
    paginas: dict[str, Pagina] = {}
    modelos: set[str] = set()
    for ws in wb.worksheets:
        filas = list(ws.iter_rows(values_only=True))
        cab_idx = next((i for i, r in enumerate(filas) if r and _txt(r[0]) == "Nº"), None)
        if cab_idx is None:
            continue  # hoja sin diseño (portada, notas…)
        cab = [_txt(c) for c in filas[cab_idx]]

        def col(nombre: str) -> int | None:
            for i, c in enumerate(cab):
                if c.lower().startswith(nombre.lower()):
                    return i
            return None

        ci = {k: col(k) for k in ("Nº", "Posic", "Lon", "Tipo", "Descrip", "Validaci", "Contenido")}
        if any(ci[k] is None for k in ("Nº", "Posic", "Lon", "Tipo", "Descrip")):
            raise ValueError(f"Hoja {ws.title}: columnas inesperadas {cab}")

        def celda(r: tuple, k: str) -> Any:
            i = ci[k]
            return r[i] if i is not None and i < len(r) else None

        campos: list[Campo] = []
        total: int | None = None
        notas: list[str] = []
        fin = False
        for r in filas[cab_idx + 1 :]:
            if not r or all(c is None for c in r):
                continue
            primero = _txt(r[0])
            if fin:
                notas.append(" | ".join(_txt(c) for c in r if c is not None))
                continue
            if primero.upper() == "TOTAL":
                total = next((n for n in (_int(c) for c in r[1:]) if n is not None), None)
                fin = True
                continue
            desc = _txt(celda(r, "Descrip"))
            tipo = _txt(celda(r, "Tipo"))
            num = _int(r[0])
            if num is None:
                # El Excel numera con fórmulas (=A6+1). Si no trae el valor calculado, se deduce.
                if not (desc and tipo):
                    continue
                num = campos[-1].num + 1 if campos else 1
            lon = _int(celda(r, "Lon"))
            previo = campos[-1] if campos else None
            esperada = 1 if previo is None else (
                previo.pos + previo.lon if previo.pos is not None and previo.lon is not None else None
            )
            pos = _int(celda(r, "Posic"))
            if pos is None:
                pos = esperada
            elif esperada is not None and pos != esperada:
                raise ValueError(f"Hoja {ws.title}: el campo {num} empieza en {pos}, se esperaba {esperada}.")
            val = _txt(celda(r, "Validaci"))
            cont = _txt(celda(r, "Contenido"))
            constante = None
            for t in (cont, val):
                m = _RE_CONSTANTE.match(t)
                if m:
                    constante = m.group(1)
                    break
            m = _RE_CASILLA.search(desc)
            campos.append(
                Campo(
                    num=num,
                    pos=pos,
                    lon=lon,
                    tipo=tipo,
                    descripcion=desc,
                    validacion=val,
                    contenido=cont,
                    casilla=m.group(1) if m else None,
                    constante=constante,
                )
            )
        if not campos:
            raise ValueError(f"Hoja {ws.title}: sin campos.")
        codigo = _codigo_pagina(campos)
        if codigo is None:
            raise ValueError(f"Hoja {ws.title}: no se identifica la página (falta el campo 'Página').")
        if codigo in paginas:
            raise ValueError(f"Hoja {ws.title}: la página {codigo} está repetida.")
        if (m := _modelo(campos)) is not None:
            modelos.add(m)
        ultimo = campos[-1]
        if total is None and codigo != "00" and ultimo.pos is not None and ultimo.lon is not None:
            total = ultimo.pos + ultimo.lon - 1
        paginas[codigo] = Pagina(ws.title, codigo, campos, total, notas)
    if len(modelos) != 1:
        raise ValueError(f"No se reconoce el modelo del Excel (encontrados: {sorted(modelos) or 'ninguno'}).")
    detectado = modelos.pop()
    if modelo and detectado != modelo:
        raise ValueError(f"El Excel es del modelo {detectado}, no del {modelo}.")
    faltan = PAGINAS_REQUERIDAS.get(detectado, {"00", "01"}) - set(paginas)
    if faltan:
        raise ValueError(f"El Excel del modelo {detectado} no contiene las páginas esperadas: {sorted(faltan)}.")
    return Diseno(path.name, sha256_de(path), paginas, detectado)


# --- Manifiestos y selección por periodo ---


@dataclass
class EntradaManifiesto:
    archivo: str
    desde: str  # trimestre desde el que aplica, p. ej. '2026-1T'
    actualizado: str | None = None  # fecha de actualización publicada por la AEAT (ISO)
    sha256: str | None = None
    fuente: str | None = None
    instalado: str | None = None  # fecha de instalación (solo diseños del usuario)
    modelo: str = "303"

    @property
    def trimestre_desde(self) -> Trimestre:
        return Trimestre.parse(self.desde)


def leer_manifiesto(directorio: Path) -> list[EntradaManifiesto]:
    p = Path(directorio) / "manifest.json"
    if not p.exists():
        return []
    return [EntradaManifiesto(**e) for e in json.loads(p.read_text(encoding="utf-8"))]


def escribir_manifiesto(directorio: Path, entradas: list[EntradaManifiesto]) -> None:
    p = Path(directorio) / "manifest.json"
    p.write_text(json.dumps([asdict(e) for e in entradas], ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


@dataclass
class DisenoDisponible:
    entrada: EntradaManifiesto
    ruta: Path
    origen: str  # 'paquete' | 'usuario'


def disponibles(dir_usuario: Path | None, modelo: str | None = None) -> list[DisenoDisponible]:
    res = [DisenoDisponible(e, DIR_EMPAQUETADO / e.archivo, "paquete") for e in leer_manifiesto(DIR_EMPAQUETADO)]
    if dir_usuario:
        res += [DisenoDisponible(e, Path(dir_usuario) / e.archivo, "usuario") for e in leer_manifiesto(dir_usuario)]
    return [d for d in res if modelo is None or d.entrada.modelo == modelo]


def seleccionar(t: Trimestre, dir_usuario: Path | None, modelo: str = "303") -> DisenoDisponible:
    """El diseño aplicable a un periodo: el de inicio más reciente que no sea posterior
    al periodo; a igualdad, el actualizado más tarde (y el del usuario sobre el del paquete).
    Para el 390 (anual) se usa el 4T del ejercicio como periodo de referencia."""
    candidatos = [d for d in disponibles(dir_usuario, modelo) if d.entrada.trimestre_desde <= t]
    if not candidatos:
        periodo = t.ejercicio if modelo == "390" else t
        raise LookupError(
            f"No hay ningún diseño de registro del modelo {modelo} instalado que aplique a {periodo}. "
            f"Ejecuta comprobar_diseno_{modelo} para descargar el vigente."
        )

    def orden(d: DisenoDisponible) -> tuple:
        return (d.entrada.trimestre_desde, d.entrada.actualizado or "", d.origen == "usuario")

    return max(candidatos, key=orden)


@lru_cache(maxsize=16)
def _cargar(ruta: str, sha: str) -> Diseno:
    return leer_xlsx(Path(ruta))


def cargar(d: DisenoDisponible) -> Diseno:
    sha = sha256_de(d.ruta)
    if d.entrada.sha256 and d.entrada.sha256 != sha:
        raise ValueError(f"El fichero {d.ruta} no coincide con el hash del manifiesto: puede estar alterado.")
    return _cargar(str(d.ruta), sha)


# --- Comparación entre versiones ---


def _clave(c: Campo) -> str:
    if c.casilla:
        return f"[{c.casilla}]"
    return re.sub(r"\s+", " ", c.descripcion.lower())


def diferencias(anterior: Diseno, nuevo: Diseno) -> dict[str, Any]:
    """Diferencias por página: campos añadidos, eliminados y modificados."""
    res: dict[str, Any] = {"paginas_nuevas": [], "paginas_eliminadas": [], "paginas": {}}
    for cod in sorted(set(nuevo.paginas) - set(anterior.paginas)):
        res["paginas_nuevas"].append(cod)
    for cod in sorted(set(anterior.paginas) - set(nuevo.paginas)):
        res["paginas_eliminadas"].append(cod)
    for cod in sorted(set(anterior.paginas) & set(nuevo.paginas)):
        a = {_clave(c): c for c in anterior.paginas[cod].campos}
        n = {_clave(c): c for c in nuevo.paginas[cod].campos}
        anadidos = [n[k].descripcion for k in n if k not in a]
        eliminados = [a[k].descripcion for k in a if k not in n]
        cambiados = []
        for k in a.keys() & n.keys():
            ca, cn = a[k], n[k]
            cambios = {
                attr: (getattr(ca, attr), getattr(cn, attr))
                for attr in ("pos", "lon", "tipo", "constante")
                if getattr(ca, attr) != getattr(cn, attr)
            }
            if cambios:
                cambiados.append({"campo": cn.descripcion, "cambios": cambios})
        tot = (anterior.paginas[cod].total, nuevo.paginas[cod].total)
        if anadidos or eliminados or cambiados or tot[0] != tot[1]:
            res["paginas"][cod] = {
                "anadidos": anadidos,
                "eliminados": eliminados,
                "modificados": cambiados,
                "longitud": {"antes": tot[0], "despues": tot[1]},
            }
    res["iguales"] = not (res["paginas_nuevas"] or res["paginas_eliminadas"] or res["paginas"])
    return res
